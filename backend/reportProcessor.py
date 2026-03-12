import framework.queryparams
import framework.postgresmodel
import re
import os
import sys
import json
import asyncio
import reportUtils
import Helpers
import datetime
import httpx
import dateutil.parser
import requests
import requests.auth
import keyCloakManager
import graphGenerator
import pptx.util
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from pptx import Presentation
import pptx.chart.data
import pandas
import numpy


# Todo:- Need to update queries for last scan time instead of "now" have to make in dashboard first

def getCookie(realm_name, clientid, role):
    keycloak_ins = keyCloakManager.KeyClockSecret()
    status, resp = keycloak_ins.create_client(clientid, realm_name, True, role)
    authKey = requests.auth._basic_auth_str(resp['clientid'], resp['clientsecret'])
    with httpx.Client(verify=False) as client:
        resp = client.get(f"https://{realm_name}.mycybercns.com/api/login", headers={"Authorization": authKey},
                          allow_redirects=False)
        if int(resp.status_code / 100) not in [2, 3]:
            return False, resp.text
        cookie = resp.cookies.get('framework')
        return True, cookie


class ReportGenerator():
    def __init__(self):
        self.reportpath = "/data/agents/runtimereports"
        if not os.path.exists(self.reportpath):
            os.makedirs(self.reportpath)
        self.reportTemplatePath = "/data/ReportTpl"
        self.ins = framework.postgresmodel.PostgresModel()
        self.ins.Config.collection_name = 'test'

    async def _getWidgetsData(self):
        widgetData = {}
        reportWidgetPath = self.reportTemplatePath + "/PptxTpl/Widgets"
        for file in os.listdir(reportWidgetPath):
            if file.endswith(".json"):
                widgetData[file.replace(".json", "")] = json.load(open(os.path.join(reportWidgetPath, file)))
        return widgetData

    async def _getOfficeData(self, reportType):
        widgetData = {}
        reportType = "/DocxTpl" if reportType == "docx" else "/XlsxTpl"
        templatesPath = self.reportTemplatePath + reportType
        for file in os.listdir(templatesPath):
            if file.endswith(".json"):
                docxData = json.load(open(os.path.join(templatesPath, file)))
                widgetData[docxData["id"]] = docxData
        return widgetData

    async def getReportsList(self):
        finalData = []
        for section in json.loads(open(self.reportTemplatePath + "/ReportList.json").read()):
            sectionDict = {"Section": section.get("Section", ""), "Reports": []}
            for report in section.get("Reports", []):
                subReports = []
                for reportType in report.get("reports", []):
                    subReports.append({"id": reportType.get("id", ""),
                                       "reportType": reportType.get("reportType", "")})
                sectionDict["Reports"].append({"title": report.get("title", ""),
                                               "reports": subReports,
                                               "description": report.get("description", "")})
            finalData.append(sectionDict)
        return finalData

    async def _transform(self, report, extracteddata):
        for transformDetails in report.get("transform", []):
            data = extracteddata.get(transformDetails["key"])
            transforms = transformDetails["transforms"]
            transformeddata = pandas.DataFrame(data)
            if transformeddata.empty:
                extracteddata[transformDetails["resp_key"]] = []
                continue
            for transform in transforms:
                transformType = transform["type"]
                if transformType == "filter":
                    for filterstring in transform["filterstrings"]:
                        if filterstring['key'] not in transformeddata.columns:
                            transformeddata[filterstring['key']] = ""
                        filterString = "transformeddata[transformeddata['{key}'] {operation} {value}]".format(
                            **filterstring)
                        transformeddata = eval(filterString)
                elif transformType == "evaluation":
                    exec(transform["execution"])
                elif transformType == "merge":
                    df = pandas.DataFrame(extracteddata.get(transform["table"]))
                    transformeddata = transformeddata.merge(df, right_on=transform["right_on"],
                                                            left_on=transform["left_on"], how='left')
                elif transformType == "lambda":
                    # todo:- need to get default value from template
                    if transform['column'] not in transformeddata.columns:
                        transformeddata[transform['column']] = ""
                    # todo:- based on dtype we have to do fillna
                    transformeddata[transform['column']].fillna('', inplace=True)
                    transformeddata[transform['resp_column']] = eval(
                        "transformeddata['{column}'].apply(lambda x: {lambdafnc})".format(**transform))
                elif transformType == "append":
                    keys = []
                    for key in transform['keys']:
                        if key in extracteddata:
                            keys.append(key)
                    transformeddata = transformeddata.append([pandas.DataFrame(extracteddata[key]) for key in keys],
                                                             ignore_index=True)
                elif transformType == "group":
                    columns = transform["groupby"] + transform.get("valuecolumns", [])
                    op = transform.get("operation", "count")
                    if op == 'count':
                        for column in columns:
                            if column not in transformeddata:
                                transformeddata[column] = numpy.NaN
                        transformeddata = transformeddata[columns]
                        transformeddata['groupCount'] = 1
                    transformeddata = transformeddata.groupby(by=transform["groupby"])
                    transformeddata = eval("transformeddata." + op + "()")
                    transformeddata = transformeddata.reset_index()
                elif transformType == "convert":
                    replacedData = {}
                    for replace, orig in transform.get('keys').items():
                        value = transformeddata.loc[transformeddata[transform['basekey']] == orig, 'groupCount'].values
                        replacedData[replace] = value[0] if value else 0
                    transformeddata = pandas.DataFrame([replacedData])
                elif transformType == "empty":
                    for key in transform["columns"]:
                        if key not in transformeddata.columns:
                            transformeddata[key] = numpy.NaN
                elif transformType == "fillna":
                    if transform.get('columns'):
                        for column, defValue in transform.get('columns').items():
                            if column in transformeddata.columns:
                                transformeddata[column] = eval("transformeddata[column].fillna(defValue)")
                            else:
                                transformeddata[column] = defValue
                    else:
                        transformeddata = eval("transformeddata.fillna(transform.get('default', "''"))")
            extracteddata[transformDetails["resp_key"]] = transformeddata.to_dict(orient='records')
        return extracteddata

    async def _extractData(self, companyid, report):
        extracteddata = {"companyid": companyid}
        for extract in report.get("extract", []):
            query = json.loads(extract.get("query", ""))
            if not query.get("query", {}).get("bool", {}).get("must"):
                query["query"]["bool"]["must"] = []
            query["query"]["bool"]["must"].append({"match": {"companyRef.id.keyword": companyid}})
            if "indexAppend" in extract:
                self.ins.Config.collection_name = 'test_timeseries'
            else:
                self.ins.Config.collection_name = 'test'
            index = await self.ins.collection_name()
            if "aggregation" in extract:
                client = await self.ins.client()
                extractdata = await client.search(index=index, body=query)
                extracteddata[extract["respKey"]] = eval(
                    str(extractdata) + ".get('aggregations', {})" + extract.get("aggregation", ""))
            else:
                params = framework.queryparams.QueryParams()
                params.limit = extract.get("size", 10000)
                params.fields = extract.get("fields", [])
                params.skip = 0
                params.q = json.dumps(query)
                extracteddata[extract["respKey"]] = (await self.ins.get_all(params)).get("data", [])
        try:
            extracteddata = await self._transform(report, extracteddata)
        except:
            extracteddata = extracteddata
        self.ins.Config.collection_name = 'test'
        return extracteddata

    async def deletePptxShape(self, shape):
        sp = shape._sp
        sp.getparent().remove(sp)

    async def insertPptxText(self, shape, data):
        if isinstance(data, int) or isinstance(data, float):
            if data >= 1000:
                data = round(data / 1000, 1)
                data = str(int(data)) if data.is_integer() else str(data)
                data = data + "k"
            else:
                data = str(data)
        if shape.has_text_frame:
            shape.text_frame.paragraphs[0].runs[0].text = data

    async def insertPptxChart(self, shape, data):
        # datamodel = {"labels": [], "values": {"series": []}}
        if shape.has_chart:
            chart_data = pptx.chart.data.CategoryChartData()
            try:
                chart_data.categories = data["labels"]
                for series, value in data["values"].items():
                    chart_data.add_series(series, value)
                shape.chart.replace_data(chart_data)
            except:
                chart_data.categories = ["No Data"]
                chart_data.add_series("No Data", [0])
                shape.chart.replace_data(chart_data)

    async def insertPptxTable(self, shape, data, length):
        col = 2
        for column in data[:length]:
            row = 0
            for k, v in column.items():
                cell = shape.table.cell(col, row)
                cell.text = str(v)
                row += 1
            col += 1

    async def letterGrade(self, riskScore):
        grade = "-"
        if isinstance(riskScore, int) or isinstance(riskScore, float):
            if 0 < riskScore < 40:
                grade = "A"
            elif 40 <= riskScore < 45:
                grade = "B"
            elif 45 <= riskScore < 60:
                grade = "C"
            elif 60 <= riskScore < 75:
                grade = "D"
            elif 75 <= riskScore < 90:
                grade = "E"
            elif 90 <= riskScore <= 100:
                grade = "F"
        return grade

    async def lastScanTime(self, companyid):
        respscantime={}
        scantimes = {
            "AssetInventoryScan": {
                "fields": "lastscannedtime",
                "sort": [{"lastscannedtime": {"order": "desc"}}],
                "query": {"query": {"bool": {"must": [{"exists": {"field": "lastscannedtime"}}]}}}},
            "ActiveDirectoryAssetInventoryScan": {
                "fields": "u",
                "sort": [{"u": {"order": "desc"}}],
                "query": {"query": {"bool": {"must": [{"exists": {"field": "object_type"}}]}}}}}
        for resp, query in scantimes.items():
            params = framework.queryparams.QueryParams()
            params.limit = 1
            params.fields = [query["fields"]]
            params.skip = 0
            params.sort = json.dumps(query["sort"])
            params.q = json.dumps(query["query"])
            querydata=(await self.ins.get_all(params)).get("data", [])
            resptime = querydata[0].get(query["fields"], "") if querydata else ""
            if resptime:
                try:
                    respscantime[resp] = str(dateutil.parser.parse(resptime)).split(".")[0]
                except:
                    respscantime[resp] = resptime
        return {k:v for k,v in respscantime.items() if v}

    async def _checkDocxData(self, transformeddata):
        datacount = 0
        for key, val in transformeddata.items():
            if isinstance(val, list):
                datacount += len(val)
        return datacount

    async def generateBulkReports(self, cookie, companyid, reportid=[], emails=[]):
        tmpDir = f"{companyid}_{datetime.datetime.utcnow().strftime('%Y_%m_%d_%H_%M')}"
        temppath = os.path.join(self.reportpath, tmpDir)
        reports = await self.getReportsList()
        if "*" in reportid or not reportid:
            reportid = []
            for record in reports:
                for repo in record["Reports"]:
                    reportid.extend([r['id'] for r in repo["reports"]])
        else:
            reports_bulk = []
            title_link = {}
            reportid_link = []
            for record in reports:
                for repo in record["Reports"]:
                    reportid_link.extend([r['id'] for r in repo["reports"]])
            for rec in reports:
                for report in rec['Reports']:
                    title_link[report['title']] = [rec['id'] for rec in report['reports']]
            for report in reportid:
                if report in title_link:
                    reports_bulk.extend(title_link[report])
                elif report in reportid_link:
                    reports_bulk.append(report)
            reportid = list(set(reports_bulk))
        resp = 'Error in report generation'
        for report_id in reportid:
            status, resp = await self.generateReport(companyid, report_id, temppath)
        files = os.listdir(temppath)
        if not files:
            return False, resp
        status, out, err = Helpers.execute("which 7z")
        if not status:
            compressionCommand = "7z a -r -t7z"
            compressionExtension = ".7z"
            compressedFile = tmpDir + ".7z"
        else:
            compressionCommand = "zip -r"
            compressionExtension = ".zip"
            compressedFile = tmpDir + ".zip"
        status, out, err = Helpers.execute(
            f"cd {self.reportpath};{compressionCommand} {compressedFile} {os.path.basename(tmpDir)}")
        Helpers.execute(f"rm -rf {temppath}")
        # todo:- need to send email if applicable
        if len(emails):

            ins = framework.postgresmodel.PostgresModel()
            ins.Config.collection_name = 'test'
            client = await ins.client()
            companyData = await client.get(index=f"test_{framework.ctx['tenant']}",
                                           id=companyid)
            companyData = {**companyData["_source"],
                           **{"_id": companyData['_id']}}
            # {"query": {"bool": {"must": [{"match": {"sourceCompanyId.keyword": "*"}}, {"match": {"destCompanyId.keyword": "*"}}]}}}
            query = Helpers.buildElsQuery(must={"sourceCompanyId.keyword": companyid,
                                                "integrationName.keyword": "Email"})
            params = framework.queryparams.QueryParams()
            params.q = json.dumps(query)
            params.limit = 1
            resp = await ins.get_all(params)
            if len(resp["data"]) == 0:
                query = Helpers.buildElsQuery(must={"sourceCompanyId.keyword": "*",
                                                    "integrationName.keyword": "Email",
                                                    "destCompanyId.keyword": "*"})
                params = framework.queryparams.QueryParams()
                params.q = json.dumps(query)
                params.limit = 1
                resp = await ins.get_all(params)
            if len(resp["data"]) > 0:
                integrationId = resp['data'][0]['credentialId']
                resp = await client.get(index=f"test_{framework.ctx['tenant']}", id=integrationId)
                credentials = {**resp["_source"], **{"_id": resp['_id']}}
                emailData = {"integrationName": "Email", "integrationId": integrationId,
                             "params": {"action": {}, "params": {"requestparams": {}}}}
                emailData["params"]["action"] = {"destination": "sendemail", "verb": "METHOD", "name": "sendEmail"}
                # todo:- need to get companyname
                emailData["params"]["params"]["requestparams"] = {
                    "summary": "Reports Generated For Company %s" % companyData["name"],
                    "initialDescription": "Reports Generated For Company %s" % companyData["name"],
                    "toEmail": ",".join(emails), "attachFiles": [temppath + compressionExtension], "isHtml": False}
                emailData["params"]["params"]["requestparams"]["fromuser"] = credentials["fromuser"]
                resp = httpx.request('POST',
                                     f"https://{framework.ctx['tenant']}.mycybercns.com/api/integrations/executeAction",
                                     json=emailData,
                                     cookies={"framework": cookie},
                                     headers={"content-type": "application/json"}, timeout=180)
        return True, temppath.replace("/opt", "") + compressionExtension

    async def generateReport(self, companyid, reportid, temppath=""):
        companyParams = {"companyid": companyid,
                         "companyName": (
                             await (await self.ins.client()).get(index=await self.ins.collection_name(), id=companyid))[
                             '_source']['name']}
        if not temppath:
            report_path = os.path.join(self.reportpath, companyid)
        else:
            report_path = temppath
        if not os.path.exists(report_path):
            os.makedirs(report_path)
        templatedict = {}
        for section in json.loads(open(self.reportTemplatePath + "/ReportList.json").read()):
            for report in section["Reports"]:
                for subReports in report["reports"]:
                    templatedict[subReports["id"]] = {
                        "title": report["title"] + " - " + re.sub(r"[^a-zA-Z0-9]+", " ", companyParams["companyName"]),
                        **subReports}
        reportType = templatedict.get(reportid, {}).get("reportType", "")
        if reportType == "pptx":
            return await self.generatePptx(companyParams, templatedict.get(reportid, {}), report_path)
        if reportType == "docx":
            return await self.generateDocx(companyParams, templatedict.get(reportid, {}), report_path)
        if reportType == "xlsx":
            return await self.generateXlsx(companyParams, templatedict.get(reportid, {}), report_path)

    async def generatePptx(self, companyParams, templatedict, report_path):
        companyid = companyParams.get("companyid")
        companyName = companyParams.get("companyName")
        widgetdata = await self._getWidgetsData()
        prs = Presentation(
            os.path.join(self.reportTemplatePath, "PptxTpl", templatedict.get("templateFile", "")))
        slidelen = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.name == "companyName":
                    await self.insertPptxText(shape, companyName)
                if shape.name in widgetdata:
                    shapeType = widgetdata.get(shape.name).get("shapeType", "")
                    extractedData = await self._extractData(companyid, widgetdata.get(shape.name))
                    externalCall = widgetdata.get(shape.name).get("externalCall", {})
                    if externalCall:
                        success = await eval("reportUtils.PptxExternalCalls().%s" % externalCall["functionName"])(shape,
                                                                                                                  extractedData,
                                                                                                                  externalCall.get(
                                                                                                                      "params",
                                                                                                                      ""))
                        success = success if success else 0
                        slidelen += success
                    else:
                        if shapeType == "chart":
                            labels, values = ["No Data"], [0]
                            if isinstance(extractedData.get(widgetdata.get(shape.name).get("respKey", ""), ""), dict):
                                labels = widgetdata.get(shape.name).get("labels", [])
                                values = [
                                    extractedData.get(widgetdata.get(shape.name).get("respKey", ""), {}).get(v, {}).get(
                                        "doc_count", 0) for v in labels]
                                if not sum(values):
                                    values = [
                                        extractedData.get(widgetdata.get(shape.name).get("respKey", ""), {}).get(v,
                                                                                                                 {}).get(
                                            "value", 0) for v in labels]
                            elif isinstance(extractedData.get("resp", ""), list):
                                labels = [v.get("key", "") for v in extractedData.get("resp", "")][
                                         :widgetdata.get(shape.name).get("chartLen", 5)]
                                values = [v.get("doc_count", 0) for v in extractedData.get("resp", "")][
                                         :widgetdata.get(shape.name).get("chartLen", 5)]
                            if sum(values):
                                slidelen += 1
                            await self.insertPptxChart(shape, {"labels": labels, "values": {"count": values}})
                        elif shapeType == "text":
                            textdata = sum([extractedData.get(widgetdata.get(shape.name).get("respKey", ""), {}).get(v,
                                                                                                                     {}).get(
                                "doc_count", 0) for v in widgetdata.get(shape.name).get("labels", [])])
                            if textdata:
                                slidelen += 1
                            await self.insertPptxText(shape, textdata)
                        elif shapeType == "table":
                            tabledata = extractedData.get(widgetdata.get(shape.name).get("respKey", ""), [])
                            if len(tabledata):
                                slidelen += 1
                            await self.insertPptxTable(shape, tabledata, widgetdata.get(shape.name).get("length", 5))
                        elif shapeType == "delete":
                            await self.deletePptxShape(shape)
        reportFile = os.path.join(report_path, templatedict.get("title", "") + ".pptx")
        prs.save(reportFile)
        if slidelen:
            return True, reportFile.replace("/opt", "")
        else:
            return False, "No data for " + templatedict.get("title", "")

    async def generateDocx(self, companyParams, templatedict, report_path):
        companyid = companyParams.get("companyid")
        docxData = await self._getOfficeData("docx")
        extractedData = await self._extractData(companyid, docxData.get(templatedict.get("id")))
        extractedData["Company"] = companyParams.get("companyName")
        extractedData["scanTime"] = await self.lastScanTime(companyid)
        tpl = DocxTemplate(os.path.join(self.reportTemplatePath, "DocxTpl", templatedict.get("templateFile", "")))
        chartImages = []
        externalCall = docxData.get(templatedict.get("id", "")).get("externalCall", {})
        if externalCall:
            externalCallData, charts = await eval("reportUtils.DocxExternalCalls().%s" % externalCall["functionName"])(
                tpl, extractedData, externalCall.get("params", ""))
            extractedData.update(externalCall)
            chartImages.extend(charts)
        for charts in docxData.get(templatedict.get("id", "")).get("loadCharts", {}):
            label, values = [], []
            chartdata = extractedData.get(charts.get("respKey", ""))
            if isinstance(chartdata, dict):
                label = charts.get("labels", [])
                values = [
                    round(chartdata.get(v, {}).get("value", 0)) for v in label]
                if not sum(values):
                    values = [
                        round(chartdata.get(v, {}).get("doc_count", 0)) for v in label]
            elif isinstance(chartdata, list):
                label = [x.get("key", "") for x in chartdata]
                values = [x.get("doc_count", "") for x in chartdata]
            graphDict = {"pie": "pieChart",
                         "hbar": "horizontalBar"}
            if sum(values):
                imageFile = await eval("graphGenerator.%s" % graphDict.get(charts.get("chartType"), "pie"))(keys=label,
                                                                                                            values=values,
                                                                                                            colours=charts.get(
                                                                                                                "colours",
                                                                                                                None),
                                                                                                            title=charts.get(
                                                                                                                "title",
                                                                                                                None))

                extractedData[charts.get("respKey", "")] = {
                    "image": InlineImage(tpl, imageFile, height=Mm(charts.get("height", 100)))}
                chartImages.append(imageFile)
            else:
                extractedData[charts.get("respKey", "")] = {"image": ""}
        # print("docxData", "*" * 10, extractedData)
        tpl.render(extractedData)
        reportFile = os.path.join(report_path, templatedict.get("title", "") + ".docx")
        tpl.save(reportFile)
        if await self._checkDocxData(extractedData) or len(chartImages):
            for imageFile in chartImages:
                if os.path.exists(imageFile):
                    os.unlink(imageFile)
            return True, reportFile.replace("/opt", "")
        else:
            return False, "No data for " + templatedict.get("title", "")

    async def generateXlsx(self, companyParams, templatedict, report_path):
        txt_clr = "#000081"
        headerclr = "#DAF7A6"
        backclr = "#FFFFFF"
        companyid = companyParams.get("companyid")
        xlsxData = await self._getOfficeData("xlsx")
        extractedData = await self._extractData(companyid, xlsxData.get(templatedict.get("id")))
        externalCall = xlsxData.get(templatedict.get("id"), {}).get("externalCall")
        if externalCall:
            extractedData = await eval("reportUtils.XlsxExternalCalls().%s" % externalCall["functionName"])(
                extractedData,
                externalCall.get(
                    "params",
                    ""))
        extractedData["Company"] = companyParams.get("companyName")
        reportFile = os.path.join(report_path, templatedict.get("title", "") + ".xlsx")
        writer = pandas.ExcelWriter(reportFile, engine='xlsxwriter')
        df = pandas.DataFrame()
        df.to_excel(writer, sheet_name='Company')
        workbook = writer.book
        worksheet = writer.sheets['Company']
        worksheet.write("H4", "Prepared for", workbook.add_format({'font_color': txt_clr,
                                                                   'border': 0,
                                                                   'font_size': 30}))
        worksheet.write("H5", extractedData.get('Company', ''), workbook.add_format({'font_color': txt_clr,
                                                                                     'border': 0,
                                                                                     'font_size': 40,
                                                                                     'bold': True}))
        timeDict = (await self.lastScanTime(companyid)).get(xlsxData[templatedict.get("id")]["scantime"], "")
        if timeDict:
            worksheet.write('H7', 'Scan performed on ' + timeDict, workbook.add_format({'font_color': txt_clr,
                                                                                        'border': 0,
                                                                                        'font_size': 20}))
        # print("xlsxData", "*" * 10, extractedData)
        sheetlen = 0
        for loadDetails in xlsxData[templatedict.get("id")]["load"]:
            if len(extractedData.get(loadDetails['respKey'], [])):
                sheetlen += 1
                df = pandas.DataFrame(
                    [{k: "-" if not v and not isinstance(v, int) else v for k, v in xlsxdata.items()} for xlsxdata in
                     pandas.json_normalize(extractedData.get(loadDetails['respKey'], []), sep='.').to_dict(
                         orient='records')])
                df.fillna("-", inplace=True)
                for key, sub in xlsxData.get(templatedict.get("id")).get("subValues", {}).items():
                    if key in df.columns:
                        df[key] = df[key].map(sub)
                keymap = {}
                for index, key in enumerate(loadDetails['tableKeys']):
                    if key not in df.columns:
                        df[key] = ""
                    keymap[key] = loadDetails["headers"][index]
                df = df[loadDetails["tableKeys"]].rename(columns=keymap)
                sheetName = loadDetails.get("sheetName", "-")[0:30]
                df.to_excel(writer, sheet_name=sheetName, index=False, startcol=0)
                worksheet = writer.sheets[sheetName]
                format = workbook.add_format()
                format.set_bold()
                worksheet.set_row(0, 18, format)
                color = workbook.add_format({"bg_color": backclr,
                                             "font_color": txt_clr,
                                             "border": 1})
                color.set_align("left")
                color.set_align("vcenter")
                header_color = workbook.add_format({"bg_color": headerclr,
                                                    "font_color": txt_clr,
                                                    "border": 1,
                                                    "font_size": 14,
                                                    "bold": True})
                header_color.set_align("left")
                header_color.set_align("vcenter")
                for i, col in enumerate(df.columns.values):
                    worksheet.set_row(0, None, color)
                    worksheet.write(0, i, col, header_color)
                    worksheet.set_column(i - 1, i, len(col) + 10, color)
        writer.save()
        if sheetlen:
            return True, reportFile.replace("/opt", "")
        else:
            return False, "No data for " + templatedict.get("title", "")


async def runReportProcessor(args):
    domainname = args["domain"]
    realm_name = domainname.split(".")[0]
    status, cookie = getCookie(f"{realm_name}",
                               "reportProcessor", "")
    if not status or not cookie:
        print("Error in generating scheduled reports, unable to fetch cookie %s" % cookie)
    ins = framework.postgresmodel.PostgresModel()
    ins.Config.collection_name = 'test'
    for record in args["records"]:
        url = record["url"]
        doc = record["doc"]
        if not doc['companyid'] or "*" in doc['companyid']:
            query = Helpers.buildElsQuery(mustExpression=[{"exists": {"field": "description"}}],
                                          mustNotExpression=[{"exists": {"field": "companyRef"}}])
            params = framework.queryparams.QueryParams()
            params.q = json.dumps(query)
            params.limit = 10000
            resp = await ins.get_all(params, cyberTenant=realm_name)
            doc['companyid'] = [record['_id'] for record in resp['data']]
        for company in doc["companyid"]:
            body = {"companyid": company,
                    "email": doc["email"],
                    "reportid": doc["sublevelparams"]}
            try:
                resp = httpx.request("POST",
                                     f"https://{domainname}{url}",
                                     headers={"content-type": "application/json"},
                                     json=body,
                                     cookies={"framework": cookie}, timeout=360, verify=False)
            except Exception as e:
                print("Exception in report processing %s" % e)
    requests.get(f"https://{realm_name}.mycybercns.com/api/logout",
                 cookies={"framework": cookie})


if __name__ == "__main__":
    # args = json.loads(base64.b64decode(sys.argv[1].encode()))
    with open(sys.argv[1]) as f:
        args = json.load(f)
    asyncio.run(runReportProcessor(args))

